from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from weekly_seo_agent.manager_document_agent.models import ExtractionStatus
from weekly_seo_agent.manager_document_agent.parsers import (
    AttachmentValidationError,
    parse_attachment,
    validate_attachment,
    validate_attachment_content,
)


def test_parse_txt_csv_tsv():
    txt = parse_attachment("note.txt", b"Line one\nLine two")
    csv = parse_attachment("table.csv", b"col1,col2\nA,B")
    tsv = parse_attachment("table.tsv", b"col1\tcol2\nA\tB")

    assert txt.extraction_status == ExtractionStatus.OK
    assert "Line one" in txt.extracted_text

    assert csv.extraction_status == ExtractionStatus.OK
    assert "col1\tcol2" in csv.extracted_text

    assert tsv.extraction_status == ExtractionStatus.OK
    assert "A\tB" in tsv.extracted_text


def test_parse_docx_and_xlsx():
    docx_buffer = BytesIO()
    docx = DocxDocument()
    docx.add_paragraph("Executive summary")
    docx.save(docx_buffer)
    parsed_docx = parse_attachment("report.docx", docx_buffer.getvalue())

    xlsx_buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Metric"
    sheet["B1"] = "Value"
    sheet["A2"] = "Clicks"
    sheet["B2"] = 42
    workbook.save(xlsx_buffer)
    parsed_xlsx = parse_attachment("data.xlsx", xlsx_buffer.getvalue())

    assert parsed_docx.extraction_status == ExtractionStatus.OK
    assert "Executive summary" in parsed_docx.extracted_text

    assert parsed_xlsx.extraction_status == ExtractionStatus.OK
    assert "# Sheet: Data" in parsed_xlsx.extracted_text
    assert "Clicks\t42" in parsed_xlsx.extracted_text


def test_parse_pdf_partial_when_no_text_layer():
    pdf_buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(pdf_buffer)

    parsed_pdf = parse_attachment("scan.pdf", pdf_buffer.getvalue())

    assert parsed_pdf.extraction_status == ExtractionStatus.PARTIAL
    assert parsed_pdf.extracted_text == ""
    assert "No extractable text" in parsed_pdf.extraction_error


def test_validate_attachment_rejects_invalid_inputs():
    with pytest.raises(AttachmentValidationError):
        validate_attachment("malware.exe", "application/octet-stream", 128)

    with pytest.raises(AttachmentValidationError):
        validate_attachment("note.txt", "text/plain", 0)

    with pytest.raises(AttachmentValidationError):
        validate_attachment("report.pdf", "text/plain", 1024)


def test_validate_attachment_content_rejects_signature_mismatch():
    with pytest.raises(AttachmentValidationError):
        validate_attachment_content("file.pdf", b"not a pdf")

    with pytest.raises(AttachmentValidationError):
        validate_attachment_content("file.docx", b"not-a-zip")

    with pytest.raises(AttachmentValidationError):
        validate_attachment_content("file.txt", b"hello\x00world")
